"""Generate the third-party claim-fidelity test corpus.

Three variants of the same short technical note, differing in exactly one number:

  A_gross_diameter_as_radius.docx  buoyancy 4x too high, because the volume was computed with
                                   the diameter where the radius belongs - a real bug pattern,
                                   and the same class of normalisation error this project made
                                   and documented in its own failure report
  B_subtle_wrong_g                  buoyancy 0.0069 N too high because g = 9.81 was used
                                   (69x the project's own tolerance, invisible to reading)
  C_control_correct.docx           all values correct

Each note states the geometry, the fluid density, the displaced volume, the mass and the
resulting forces. Variant A is therefore checkable by arithmetic from the note's own numbers:
rho*g*V contradicts the stated buoyancy. That matters for the experiment, because it removes the
excuse that the reader could not have known the truth.
"""

from __future__ import annotations

import math
import pathlib
import sys

from docx import Document
from docx.shared import Pt

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

OUT = pathlib.Path(
    r"C:\Users\xiangming\ANRI-Aerospace\Benchmarks\EXT-001_ThirdParty_Claim_Fidelity\corpus"
)
OUT.mkdir(parents=True, exist_ok=True)

D, L, M = 0.090, 0.323, 2.030
RHO = 1000.0
G_TRUE = 9.80665

V = math.pi * (D / 2) ** 2 * L
W = M * G_TRUE
FB_TRUE = RHO * G_TRUE * V
FB_A = RHO * G_TRUE * math.pi * D**2 * L          # diameter used as radius
FB_B = RHO * 9.81 * V                              # wrong g


def build(path: pathlib.Path, f_buoy: float, *, note_tag: str) -> None:
    net = f_buoy - W
    ratio = f_buoy / W

    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    doc.add_heading("Technical Note TN-2026-04", level=0)
    doc.add_heading(
        "Hydrostatic Load Estimate for a Cylindrical Body During Vertical Water Exit",
        level=1,
    )
    doc.add_paragraph(
        "Prepared for internal design review. This note records the hydrostatic baseline "
        "used to size the vertical load case during water exit. It is a working estimate "
        "and is not a validation against measurement."
    )

    doc.add_heading("1. Body and fluid data", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Quantity"
    hdr[1].text = "Symbol"
    hdr[2].text = "Value"
    rows = [
        ("Body diameter", "D", f"{D:.3f} m"),
        ("Body length", "L", f"{L:.3f} m"),
        ("Body mass", "m", f"{M:.3f} kg"),
        ("Water density", "rho", f"{RHO:.0f} kg/m3"),
        ("Gravitational acceleration", "g", f"{G_TRUE:.5f} m/s2"),
    ]
    for q, s, v in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = q, s, v

    doc.add_heading("2. Derived quantities", level=2)
    doc.add_paragraph(
        f"The displaced volume at full submersion follows from the cylindrical geometry, "
        f"giving V = {V:.9e} m3."
    )
    doc.add_paragraph(
        f"The weight of the body is W = m g = {W:.6f} N."
    )
    doc.add_paragraph(
        f"The hydrostatic buoyancy force at full submersion is F_b = rho g V = "
        f"{f_buoy:.6f} N."
    )
    doc.add_paragraph(
        f"Consequently the net upward hydrostatic force on a fully submerged body is "
        f"{net:.4f} N, and the buoyancy-to-weight ratio is {ratio:.3f}."
    )

    doc.add_heading("3. Load case for the exit transient", level=2)
    doc.add_paragraph(
        f"During the vertical exit transient the submerged volume decreases monotonically "
        f"from V to zero, so the buoyancy falls from {f_buoy:.4f} N to zero while the weight "
        f"remains {W:.4f} N. The vertical load case is therefore taken as the fully submerged "
        f"condition, {f_buoy:.4f} N, applied together with the body weight."
    )

    doc.add_heading("4. Limitations", level=2)
    doc.add_paragraph(
        "This note treats the free surface as flat and excludes surface tension, drag, added "
        "mass and any free-surface depression effect. Attitude is assumed vertical. No "
        "measurement data are used and no uncertainty estimate is provided."
    )

    doc.save(path)
    print(f"  wrote {path.name}  (F_b = {f_buoy:.6f} N, ratio = {ratio:.3f})")


print("=== generating corpus ===")
build(OUT / "A_gross_diameter_as_radius.docx", FB_A, note_tag="A")
build(OUT / "B_subtle_wrong_g.docx", FB_B, note_tag="B")
build(OUT / "C_control_correct.docx", FB_TRUE, note_tag="C")

print()
print("=== reference values for scoring ===")
print(f"  correct F_b  = {FB_TRUE:.6f} N")
print(f"  variant A    = {FB_A:.6f} N   (ratio {FB_A / FB_TRUE:.3f}, net {FB_A - W:.3f} N)")
print(f"  variant B    = {FB_B:.6f} N   (excess {(FB_B - FB_TRUE):.6f} N = "
      f"{(FB_B - FB_TRUE) / 1e-4:.0f} tolerances)")
print(f"  weight       = {W:.6f} N")
